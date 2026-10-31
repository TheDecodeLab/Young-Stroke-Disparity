#!/usr/bin/env python3
"""Regenerate the revised manuscript Figure 1 from frozen Step68 aggregate.

Only four directly named modality groups are shown, separately by source P.
The CPT14 flags measure recorded code presence, not performed scans.
This is a LOCAL-ONLY exhibit builder, not a source clinical-data scan.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"results/phase1/full_reconstruction_private/step68_procedure_windows_private/step68_aggregate_LOCAL_ONLY.json"
OUTPUT=ROOT/"results/phase1/full_reconstruction_private/figure1_v18_from_step68_PRIVATE.png"
COHORT="uniform_18_50"
MODALITIES=(("CT_HEAD","Head CT"),("MRI_BRAIN","Brain MRI"),
            ("CTA_HEAD","Head CTA"),("MRA_HEAD","Head MRA"))
ACK="I_CONFIRM_INTERNAL_REVIEW_ONLY"

def get_data(obj):
    if obj.get("stage") != "STEP68_SINGLE_SCAN_CPT14_WINDOWS_AND_VERSIONED_SYSTEM_DISCOVERY_V2":
        raise ValueError("Expected completed frozen Step68 aggregate")
    records=obj.get("patient_cohort_aggregates",{})
    cohort_row=records.get(f"COHORT|{COHORT}")
    if not cohort_row or int(cohort_row["n"])<=0:
        raise ValueError("Missing primary cohort denominator")
    groups={}
    for p in (0,1):
        pk=f"P|{COHORT}|P={p}"
        if pk not in records or int(records[pk]["n"])<=0:
            raise ValueError("Source P-group denominator mismatch")
        groups[p]=int(records[pk]["n"])
    if sum(groups.values())!=int(cohort_row["n"]):
        raise ValueError("Source P groups fail primary cohort reconciliation")
    ans={}
    for p in (0,1):
        row=[]
        for code,label in MODALITIES:
            k=f"MODALITY_P|{COHORT}|P={p}|{code}"
            if k not in records:
                raise ValueError("Missing modality in frozen Step68 aggregate: "+k)
            rec=records[k]
            n,base,gain=(int(rec.get(x,0)) for x in ("n","imp","new_day0_other"))
            if n!=groups[p] or not (0<=base<=n and 0<=gain<=n-base):
                raise ValueError("Non-binomial modality counts or wrong cohort denominator")
            row.append({"code":code,"label":label,"P":p,
                        "index_pct":100*base/n,"same_day_pct":100*(base+gain)/n})
        ans[p]=row
    return ans

def verify_supplement(data, supplement_path):
    from docx import Document
    d=Document(str(supplement_path))
    matches=[t for t in d.tables if t.rows and t.cell(0,0).text.strip()=="Modality"
             and "I63 flag" in t.cell(0,1).text]
    if len(matches)!=1:
        raise ValueError("Expected one Supplemental Table S4")
    rows={}
    for cells in matches[0].rows[1:]:
        a=[c.text.strip() for c in cells.cells]
        if len(a)<4:
            raise ValueError("Malformed Table S4")
        rows[(a[0],a[1])]=(float(a[2]),float(a[3]))
    for p, x in data.items():
        for point in x:
            key=(point["label"],f"P={p}")
            if key not in rows:
                raise ValueError("S4 missing displayed modality "+str(key))
            vals=rows[key]
            if abs(vals[0]-point["index_pct"])>0.01001 or abs(vals[1]-point["same_day_pct"])>0.01001:
                raise ValueError("Figure data disagree with S4 at "+str(key))
    return True

def draw(data,output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    f,axes=plt.subplots(1,2,figsize=(10.5,3.9),sharex=True,sharey=True)
    fig_y=list(range(len(MODALITIES)))
    for p,ax in enumerate(axes):
        for y,point in zip(fig_y,data[p]):
            ax.plot([point["index_pct"],point["same_day_pct"]],[y,y],
                    color="#9ba2a6",linewidth=1.75,zorder=1)
            ax.scatter(point["index_pct"],y,marker="o",s=58,
                       facecolor="#293136",edgecolor="#293136",zorder=3)
            ax.scatter(point["same_day_pct"],y,marker="s",s=68,
                       facecolor="white",edgecolor="#086998",linewidth=2,zorder=4)
        ax.set_title("P=0: principal I63 not marked" if p==0 else
                     "P=1: principal I63 marked",fontsize=11,pad=8)
        ax.set_yticks(fig_y,labels=[name for _,name in MODALITIES])
        ax.grid(axis="x",alpha=0.16)
        ax.set_xlim(0,50)
        ax.set_xticks((0,10,20,30,40,50))
        ax.tick_params(axis="both",labelsize=9)
        for side in ("left","right","top"):
            ax.spines[side].set_visible(False)
    axes[0].invert_yaxis()
    f.supxlabel("Patients with a recorded CPT code (%)",fontsize=10,y=0.12)
    keys=[Line2D([0],[0],marker="o",color="none",markerfacecolor="#293136",
                  markeredgecolor="#293136",markersize=7,label="Index inpatient"),
          Line2D([0],[0],marker="s",color="none",markerfacecolor="white",
                  markeredgecolor="#086998",markeredgewidth=2,markersize=7,
                  label="Index or other same-day encounter")]
    f.legend(handles=keys,ncol=2,loc="lower center",frameon=False,fontsize=9,
             bbox_to_anchor=(0.5,-0.015))
    f.subplots_adjust(left=0.14,right=0.985,bottom=0.24,top=0.87,wspace=0.14)
    f.savefig(output,dpi=250,facecolor="white")
    plt.close(f)

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--step68",type=Path,default=SOURCE)
    ap.add_argument("--supplement",type=Path,
                    help="Optional independent S4 display-value reconciliation")
    ap.add_argument("--output",type=Path,default=OUTPUT)
    ap.add_argument("--acknowledge",required=True)
    args=ap.parse_args()
    if args.acknowledge!=ACK:
        ap.error("Explicit internal-review acknowledgement required")
    safe=(ROOT/"results/phase1/full_reconstruction_private").resolve()
    if safe not in args.output.resolve().parents:
        raise ValueError("Figure must be stored in gitignored private results")
    if args.step68.is_symlink() or not args.step68.is_file():
        raise FileNotFoundError("Frozen Step68 aggregate absent")
    if args.output.exists() or args.output.is_symlink():
        raise FileExistsError("Refusing overwrite of existing figure")
    data=get_data(json.loads(args.step68.read_text(encoding="utf8")))
    if args.supplement:
        verify_supplement(data,args.supplement)
    args.output.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
    draw(data,args.output)
    os.chmod(args.output,0o600)
    print("Figure 1 built from frozen Step68 modality aggregates (no clinical-table scan)")
    print("S4 reconciliation:", "PASS" if args.supplement else "NOT REQUESTED")
    print("Private output:",args.output)
    print("Not permitted for GitHub or external publication before applicable review.")

if __name__=="__main__":
    main()
